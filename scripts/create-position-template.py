from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
out = ROOT / 'templates'
out.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string('282C29')
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.04
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(26)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
image = p.add_run().add_picture(str(ROOT / 'assets/logo-header.png'), width=Inches(2.55))
image._inline.docPr.set('descr', 'SOPliance One Step Beyond logo')
doc.add_paragraph('Position title', 'Title')
p = doc.add_paragraph('SOPliance, Inc.  |  Position description')
p.paragraph_format.space_after = Pt(12)

def field(label, placeholder):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(label + '  ').bold = True
    p.add_run(placeholder)

field('Position reference', '[Unique reference number]')
field('Location and work arrangement', '[City, state or region] | [On site, hybrid or remote]')
field('Employment type and duration', '[Full time, part time or contract] | [Duration, if applicable]')
field('Schedule and travel', '[Working hours, shift and any travel requirements]')
field('Compensation', '[Approved range, currency and hourly or annual basis]')
field('Posted and closing dates', '[Posted date] | [Closing date or open until filled]')

def section(title, text):
    doc.add_paragraph(title, 'Heading 1')
    doc.add_paragraph(text)

def bullet(text):
    p = doc.add_paragraph(text, 'List Bullet')
    p.paragraph_format.space_after = Pt(3)

section('Role overview', '[Summarize the purpose of this position, the team or project it supports, and its main outcomes in two or three sentences.]')
doc.add_paragraph('Key responsibilities', 'Heading 1')
for text in ['[Describe a primary responsibility and the expected outcome.]', '[Describe another key responsibility or deliverable.]', '[Describe collaboration, reporting or service responsibilities.]']:
    bullet(text)
doc.add_paragraph('Required qualifications', 'Heading 1')
for text in ['[Specify relevant experience and essential skills.]', '[Specify required education, certifications or licenses, if applicable.]']:
    bullet(text)
section('Preferred qualifications', '[List useful additional experience or skills. Remove this section if none apply.]')
section('Work requirements', '[Describe role-specific work conditions, on-site attendance, equipment or other essential requirements. Remove items that do not apply.]')
doc.add_paragraph('How to apply', 'Heading 1')
p = doc.add_paragraph('Email your résumé to ')
part = p.part
rel = part.relate_to('mailto:info@sopliance.com', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
color = OxmlElement('w:color'); color.set(qn('w:val'), 'BA3D12'); props.append(color)
under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); props.append(under)
run.append(props); t = OxmlElement('w:t'); t.text = 'info@sopliance.com'; run.append(t); link.append(run); p._p.append(link)
p.add_run(' with the subject line ')
p.add_run('Application: [Position title] | [Position reference]').bold = True
p.add_run('. Include [any additional application materials required for this role].')
doc.core_properties.title = 'SOPliance Position Description Template'
doc.core_properties.subject = 'Reusable template for a specific advertised position'
doc.core_properties.author = 'SOPliance, Inc.'
doc.core_properties.keywords = 'SOPliance, careers, position, job description, template'
for tree in [doc._element, doc.styles.element]:
    for border in tree.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(out / 'SOPliance Position Template.docx')
print(out / 'SOPliance Position Template.docx')
